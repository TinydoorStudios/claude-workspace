#!/usr/bin/env python3
"""Show Status Log — a running, auto-regenerated view across every show in the
advance pipeline (both venues), for "what's the status of X at a glance."

Pulls live from advance-db (the advance_status view + a read-only join against
bookings for date-booked/location/who-booked-it) — same source, same join
pattern as the Show Status Log demo. NOT a mirror of advance-list.xlsx: that
sheet is the working/editing surface with band-detail overrides; this is a
lighter read-mostly status board.

Most columns are fully auto (overwritten fresh every run). A few have no data
source yet — Owner, Advance Deadline, Basic Advance Complete?, Info Complete?,
Feedback Survey Sent?, Notes — and are safe to hand-type into the Dropbox file
directly: a hidden Show ID column keys each row, so a re-run finds the same
show and PRESERVES whatever was typed into those manual columns rather than
blanking them. New shows just get blank manual cells; nothing is ever deleted
out from under a real show, only refreshed.

Called from run_now.py after every booking/submission (best-effort — a bug
here should never block the day-sheet/email pipeline). Can also run by hand:

    python3 status_log.py [--out PATH]
"""
import argparse
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
for _cand in (HERE.parent, HERE.parent / "app"):
    if (_cand / "advance_db.py").exists():
        sys.path.insert(0, str(_cand))
        break
import advance_db as db

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(HERE))
import fieldspec as fs
NYQUIST_DEFAULT = fs.nyquist_root()
DEFAULT_OUT = NYQUIST_DEFAULT / "Show Status Log.xlsx"

NAVY = "1A3A5C"
WHITE = "FFFFFF"
MANUAL_TINT = "FFF3CD"  # matches the amber TOUR/hand-filled convention elsewhere

# same state colors as merge_status.py's STATE_FILL, for one consistent palette
# across the whole pipeline. "responded" relabeled "Advancing In Progress" and
# "finalized" added (Brian, 2026-09-13) — a human sign-off step after the band
# responds, before the advance is really done (app.py's POST /artist/<id>/
# finalize/<show_id>). "Completed" used to mean what "Advancing In Progress"
# means now — that word moved to "finalized" instead, same wording the live
# dashboard uses for the same state, so nothing here says something different
# from that page ever again.
sys.path.insert(0, str(HERE.parent))
import status_labels as SL
# one shared label set for every surface (audit #17)
STATE_STYLE = {k: (v[0], v[2], v[3]) for k, v in SL.STATE.items()}

# (label, width, "live" | "manual")
COLS = [
    ("Event Name", 22, "live"), ("Band", 20, "live"), ("Venue", 14, "live"),
    ("Location", 11, "live"), ("Series", 16, "live"),
    ("Owner", 10, "manual"), ("Booked By", 12, "live"),
    ("Show Date", 11, "live"), ("Days Until Show", 9, "live"),
    ("Date Booked", 11, "live"), ("Advance Deadline", 12, "manual"),
    ("Advance Status", 14, "live"),
    ("Basic Advance Complete?", 11, "manual"), ("Info Complete?", 10, "manual"),
    ("Advance Drafted", 12, "live"),
    ("Follow-up Drafted", 12, "live"), ("Responded", 12, "live"),
    ("Contact Email", 22, "live"), ("Files", 6, "live"),
    ("Feedback Survey Sent?", 11, "manual"), ("Notes", 28, "manual"),
]
MANUAL_LABELS = {lbl for lbl, _w, kind in COLS if kind == "manual"}
ID_COL = len(COLS) + 1  # hidden Show ID, past every visible column


def fdate(v):
    """timestamptz -> 'YYYY-MM-DD'. Slice, don't parse — sidesteps +00 vs
    +00:00 / variable microsecond-digit fussiness; only the date is shown."""
    if v is None:
        return None
    s = str(v)
    return s[:10] if len(s) >= 10 else None


def fetch_rows(cur):
    cur.execute(
        """SELECT
             ast.show_id, ast.band, ast.venue, ast.show_series, ast.show_date,
             ast.state, ast.days_until_show,
             ast.advance_draft_created_at, ast.send_reminder_sent_at,
             ast.followup_draft_created_at, ast.responded_at, ast.contact_email,
             b.location, b.created_at AS date_booked, b.event_name, b.entered_by,
             (SELECT count(*) FROM files f WHERE f.artist_id = ast.artist_id) AS file_count
           FROM advance_status ast
           LEFT JOIN bookings b ON b.venue = ast.venue
                  AND b.event_date = ast.show_date
                  AND lower(btrim(regexp_replace(b.artist_name, '\\s+', ' ', 'g'))) =
                      lower(btrim(regexp_replace(ast.band, '\\s+', ' ', 'g')))
           ORDER BY ast.show_date ASC NULLS LAST, ast.band"""
    )
    return cur.fetchall()


class UnreadableLog(Exception):
    pass


def read_manual(path):
    """{show_id: {label: value}} preserved from the existing file. Raises
    UnreadableLog when the file exists but can't be read safely (open in
    Excel, mid-sync, damaged) — the caller must NOT overwrite it then
    (audit #19)."""
    if not path.exists():
        return {}
    if any(p.name[2:] and path.name.endswith(p.name[2:]) for p in path.parent.glob("~$*")):
        raise UnreadableLog("open in Excel")
    try:
        wb = load_workbook(path)
    except Exception as e:
        raise UnreadableLog(repr(e))
    if "Show Status" not in wb.sheetnames:
        raise UnreadableLog("no 'Show Status' tab")
    ws = wb["Show Status"]
    label_col = {ws.cell(4, c).value: c for c in range(1, len(COLS) + 1)}
    out = {}
    for r in range(5, ws.max_row + 1):
        sid = ws.cell(r, ID_COL).value
        if sid is None:
            continue
        row = {}
        for label in MANUAL_LABELS:
            c = label_col.get(label)
            if c:
                v = ws.cell(r, c).value
                if v not in (None, ""):
                    row[label] = v
        if row:
            out[int(sid)] = row
    return out


def build(rows, manual_by_id):
    wb = Workbook()
    ws = wb.active
    ws.title = "Show Status"
    n = len(COLS)
    last_letter = get_column_letter(n)

    thin = Side(style="thin", color="D9DEE5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells(f"A1:{last_letter}1")
    t = ws.cell(1, 1, "Band Advance — Show Status Log")
    t.font = Font(bold=True, color=WHITE, size=14)
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 26

    ws.merge_cells(f"A2:{last_letter}2")
    sub = ws.cell(2, 1,
        "Auto-generated from advance-db after every booking/submission. White columns "
        "refresh every run; amber columns (Owner, Advance Deadline, Basic Advance "
        "Complete?, Info Complete?, Feedback Survey Sent?, Notes) are yours to hand-type "
        "here — they're preserved across regenerations, not overwritten.")
    sub.font = Font(italic=True, color="6B7280", size=9)
    sub.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[2].height = 28

    hrow = 4
    for i, (label, width, _kind) in enumerate(COLS, start=1):
        c = ws.cell(hrow, i, label)
        c.font = Font(bold=True, color=WHITE, size=10)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[hrow].height = 30
    ws.freeze_panes = f"F{hrow + 1}"

    id_letter = get_column_letter(ID_COL)
    ws.cell(hrow, ID_COL, "Show ID")
    ws.column_dimensions[id_letter].hidden = True

    r = hrow + 1
    for row in rows:
        label, fg, txt = STATE_STYLE.get(row["state"], (row["state"] or "—", "E5E7EB", "374151"))
        preserved = manual_by_id.get(row["show_id"], {})
        live = {
            "Event Name": row["event_name"] or "—",
            "Band": row["band"],
            "Venue": row["venue"],
            "Location": row["location"] or "—",
            "Series": row["show_series"] or "—",
            "Booked By": row["entered_by"] or "—",
            "Show Date": row["show_date"].isoformat() if row["show_date"] else "",
            "Days Until Show": row["days_until_show"] if row["days_until_show"] is not None else "",
            "Date Booked": fdate(row["date_booked"]) or "",
            "Advance Status": label,
            "Advance Drafted": fdate(row["advance_draft_created_at"]) or "",
            "Follow-up Drafted": fdate(row["followup_draft_created_at"]) or "",
            "Responded": fdate(row["responded_at"]) or "",
            "Contact Email": row["contact_email"] or "—",
            "Files": row["file_count"] or 0,
        }
        for i, (clabel, _w, kind) in enumerate(COLS, start=1):
            value = live.get(clabel, preserved.get(clabel, "")) if kind == "live" else preserved.get(clabel, "")
            cell = ws.cell(r, i, value)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center" if clabel not in ("Event Name", "Band", "Contact Email", "Notes") else "left",
                vertical="center")
            if clabel == "Advance Status":
                cell.fill = PatternFill("solid", fgColor=fg)
                cell.font = Font(bold=True, color=txt, size=10)
            elif kind == "manual":
                cell.fill = PatternFill("solid", fgColor=MANUAL_TINT)
            elif r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F7FAFC")
        ws.cell(r, ID_COL, row["show_id"])
        r += 1
    last_row = max(r - 1, hrow + 1)

    tbl = Table(displayName="ShowStatus", ref=f"A{hrow}:{last_letter}{last_row}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(tbl)

    lg = wb.create_sheet("Legend")
    lg.column_dimensions["A"].width = 18
    lg.column_dimensions["B"].width = 74
    t2 = lg.cell(1, 1, "Advance Status — what each stage means")
    t2.font = Font(bold=True, size=13, color=NAVY)
    lg.merge_cells("A1:B1")
    meanings = dict(SL.MEANINGS)
    row_i = 3
    for label, desc in meanings.items():
        fg, txt = next(((st[1], st[2]) for st in STATE_STYLE.values() if st[0] == label), ("E5E7EB", "374151"))
        c = lg.cell(row_i, 1, label)
        c.fill = PatternFill("solid", fgColor=fg)
        c.font = Font(bold=True, color=txt)
        c.alignment = Alignment(horizontal="center", vertical="center")
        d = lg.cell(row_i, 2, desc)
        d.alignment = Alignment(wrap_text=True, vertical="center")
        lg.row_dimensions[row_i].height = 32
        row_i += 1
    row_i += 1
    note = lg.cell(row_i, 1, "Manual columns:")
    note.font = Font(bold=True)
    lg.merge_cells(f"A{row_i + 1}:B{row_i + 2}")
    lg.cell(row_i + 1, 1,
        "Owner / Advance Deadline / Basic Advance Complete? / Info Complete? / Feedback "
        "Survey Sent? / Notes have no data source yet — type into them directly in this "
        "file. A hidden Show ID column keys each row so the next auto-regeneration finds "
        "your entry and keeps it, instead of overwriting it blank."
    ).alignment = Alignment(wrap_text=True, vertical="top")

    return wb, last_row - hrow


FAIL_STATE = HERE.parent / "data" / "status_log_failures.json"
BACKUP_KEEP = 30


def _failures(update=None):
    import json
    try:
        data = json.loads(FAIL_STATE.read_text())
    except (OSError, ValueError):
        data = {"count": 0, "alerted": False}
    if update is not None:
        data = update
        FAIL_STATE.parent.mkdir(parents=True, exist_ok=True)
        FAIL_STATE.write_text(json.dumps(data))
    return data


def main():
    import datetime as dt
    import shutil
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()
    args.out.parent.mkdir(parents=True, exist_ok=True)

    try:
        manual_by_id = read_manual(args.out)
    except UnreadableLog as e:
        # skip, retry next run, alert after 3 in a row (audit #19)
        st = _failures()
        st["count"] = st.get("count", 0) + 1
        print(f"  ! Show Status Log not updated — existing file unreadable ({e}); "
              f"failure {st['count']} in a row", file=sys.stderr)
        if st["count"] >= 3 and not st.get("alerted"):
            try:
                import mailer
                ok, _ = mailer.alert(
                    "Show Status Log hasn't updated in 3 runs",
                    f"<p>The last {st['count']} pipeline runs couldn't read "
                    f"<code>{mailer.esc(args.out.name)}</code> ({mailer.esc(str(e))}), so it was left "
                    "untouched rather than overwritten. Close it in Excel / check for a Dropbox "
                    "conflicted copy.</p>")
                st["alerted"] = bool(ok)
            except Exception as mail_err:  # noqa: BLE001
                print(f"  ! alert failed: {mail_err!r}", file=sys.stderr)
        _failures(st)
        return

    with db.get_conn() as conn, conn.cursor() as cur:
        rows = fetch_rows(cur)

    if args.out.exists():
        bdir = args.out.parent / "_status_log_backups"
        bdir.mkdir(exist_ok=True)
        shutil.copy2(args.out, bdir / f"Show Status Log {dt.datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx")
        for old in sorted(bdir.glob("Show Status Log *.xlsx"))[:-BACKUP_KEEP]:
            old.unlink(missing_ok=True)

    wb, n_shows = build(rows, manual_by_id)
    tmp = args.out.parent / f".~tmp-{args.out.name}"
    wb.save(tmp)
    tmp.replace(args.out)
    _failures({"count": 0, "alerted": False})
    print(f"Show Status Log: {n_shows} show(s) -> {args.out} "
          f"({len(manual_by_id)} row(s) had manual cells preserved)")


if __name__ == "__main__":
    main()
