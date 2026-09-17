#!/usr/bin/env python3
"""Filed advance docs are never overwritten (2026-09-13 audit, items 1 + 2).

A show's advance doc is created ONCE, from the universal template. After
that, every pipeline pass (a booking, a band submission, "Run again")
builds what the doc WOULD say from scratch in memory, then walks the filed
doc cell by cell:

  - the pipeline has nothing for a cell (fresh == blank template)  -> skip
  - the filed doc already says the same thing                       -> skip
  - the filed doc's cell is still blank (== the blank template)     -> fill it
  - the filed doc says something different                          -> leave
    it alone and record a notice (emailed to Brian as a diff)

"Blank" means the cell still matches the pristine universal template
exactly — "FOH – " with no name, a checkbox group with nothing ticked.
Multi-paragraph cells (Engineer's FOH/Mon lines, Event Type's two checkbox
lines, the EVENT INFORMATION header's Date/Event/Venue/Location lines) are
compared paragraph by paragraph when their shape still matches the
template, so filling the Mon line never touches a hand-typed FOH line.

Per-cell provenance (review 2026-09-14, H2 — Brian's call): `filed_docs.cells`
records what the pipeline last wrote into each cell. A cell that still reads
exactly what the pipeline wrote is the pipeline's to update (a returning
band's prior-show pre-fill, a corrected monitor count, a cancelled act);
a cell that reads anything else was typed by a person and is never touched —
a differing value becomes a notice instead. "Blank" cells (still the
template) are filled as before. Whole-file sha256 is still recorded so a
hand-edited doc is flagged in the notice.

Resubmission review (Brian, 2026-09-14): once a show has a SECOND
submission (or a staff edit from the dashboard) newer than the doc's last
pipeline write, every differing filled cell is held — pipeline-owned or
hand-typed — and becomes a notice Brian decides on /doc-review (Keep doc /
Use new). Blank cells still fill. A cell with an open notice is frozen in
every later pass until he decides. "Use new" re-runs the merge with that
cell forced (only if the fresh value still matches what he saw); "Keep doc"
drops the pipeline's ownership of the cell so it's treated as hand-typed
from then on.
Provenance is keyed by ARTIST, not column (root cause B, audit 2026-09-16).
`filed_docs.columns` records which artist sat in each act column when the doc
was last written; every band-row cell is keyed `Section|a<artist_id>`. When a
new act with an earlier start reorders the bill, the filed doc's cell XML is
physically moved to the artist's new column first (hand edits included), and
the column that artist left gets the template back — so an artist's answers,
notices, "Keep doc" and "Use new" all follow the artist. Pre-2026-09-16 keys
(`Section|colN`, `Section|<band name>`) are still read as aliases. A band row
the pipeline filled goes back to the template when the fresh build has nothing
for it (a cancelled or removed act). A pipeline value somebody deleted by hand
is a notice, never a refill.

Open in Word: Dropbox never syncs Word's "~$" owner file, so the VM can't see
a doc open on a Mac. What it CAN see is the symptom — a "(… conflicted copy
…)" file beside the doc after both sides saved — and that raises a notice.
The old owner-file check stays as belt-and-braces for a doc opened on the VM
itself. Writes are atomic (temp file + os.replace) and re-check the hash right
before replacing, so a save that lands mid-run is never clobbered.

The filename follows fieldspec.event_display_name (series/event + top of bill).
When that name changes, the SAME file is renamed in place — never a second
copy. Lookups go through `filed_docs` by venue + date + event key, not by
filename sort order. Past shows are never touched.
"""
import copy
import datetime as dt
import filecmp
import hashlib
import os
import re
import shutil
import sys
from pathlib import Path
from urllib.parse import quote

HERE = Path(__file__).resolve().parent
for _c in (HERE.parent, HERE.parent / "app"):
    if (_c / "advance_db.py").exists():
        sys.path.insert(0, str(_c))
        break
sys.path.insert(0, str(HERE))
import advance_db as db
import daysheet
import fieldspec as fs

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Fallback column names when a doc's own act-name row is blank. Artist 1/2/3
# since 2026-09-15; a doc filed before that reads Opener/Direct Support/
# Headliner in the template itself, and this is only ever a fallback for a
# column whose name cell is empty, so it never relabels what's actually filed.
ACT_COL_NAMES = {1: "Artist 1", 2: "Artist 2", 3: "Artist 3"}


# ── small XML helpers ───────────────────────────────────────────────────────

def _t(s):
    return re.sub(r"\s+", " ", (s or "")).strip()


def _p_text(p_el):
    """Visible text of a paragraph, including checkbox glyphs inside content
    controls; line breaks and tabs count as whitespace so a value written as
    'a\nb' compares equal however Word stored the break."""
    out = []
    for el in p_el.iter(qn("w:t"), qn("w:br"), qn("w:tab"), qn("w:cr")):
        out.append(el.text or "" if el.tag == qn("w:t") else " ")
    return "".join(out)


def _tc_text(tc):
    return "\n".join(_p_text(p) for p in tc.iter(qn("w:p")))


def _tc_paras(tc):
    return [c for c in tc if c.tag == qn("w:p")]


def _has_table(tc):
    return any(True for _ in tc.iter(qn("w:tbl")))


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def event_key(ev):
    return re.sub(r"\s+", " ", (ev.get("name") or ev.get("series") or "").strip()).lower()


def ci_existing(folder, name):
    """The file in `folder` whose name equals `name` ignoring case, or None.
    The VM's filesystem is case-sensitive; Brian's Mac and Dropbox are not
    (2026-09-14: 'Stageplot.pdf' filed beside a hand-typed 'stageplot.pdf'
    produced a Dropbox '(Case Conflict)' copy on every run)."""
    folder = Path(folder)
    if not folder.exists():
        return None
    low = name.lower()
    for p in folder.iterdir():
        if p.name.lower() == low:
            return p
    return None


def word_lock_present(path):
    """True if a '~$' owner file sits beside this doc. Belt-and-braces only:
    Dropbox's ignore list never syncs '~$' files, so a doc open in Word on a
    Mac is invisible here (F2, audit 2026-09-16) — conflicted_copies() is the
    check that actually fires for that case."""
    path = Path(path)
    try:
        for p in path.parent.glob("~$*"):
            tail = p.name[2:]
            if tail and path.name.endswith(tail):
                return True
    except OSError:
        return False
    return False


def conflicted_copies(path):
    """Dropbox '(… conflicted copy …)' siblings of this file — what's left when
    a person saved it in Word while the pipeline wrote it too. Their edits
    live only in the copy until somebody merges them by hand."""
    path = Path(path)
    try:
        return sorted(path.parent.glob(f"{glob_escape(path.stem)} (*conflicted copy*){path.suffix}"))
    except OSError:
        return []


def glob_escape(s):
    return re.sub(r"([*?\[])", r"[\1]", s)


class _Ctx:
    """Everything one merge pass needs: the three documents, a running max
    content-control id for the filed doc, the notice list, and the per-cell
    provenance map in (`prov`, what the pipeline last wrote) and out
    (`newprov`, what the pipeline owns after this pass)."""

    def __init__(self, T, F, E, prov=None, review=False, frozen=None, force=None):
        self.T, self.F, self.E = T, F, E
        self.changed = 0
        self.notices = []
        self.review = review
        self.frozen = {k.lower() for k in (frozen or ())}
        self.force = {k.lower(): v for k, v in (force or {}).items()}
        self.settled = set()   # keys where the doc now reads the fresh value
        self.forced = set()    # keys written because Brian chose "Use new"
        self.prov = dict(prov or {})
        # case-insensitive view: legacy keys carried the band's display name
        # ("Monitors|RatBoys"); a case-only rename must not orphan them
        self.prov_ci = {k.lower(): v for k, v in self.prov.items()}
        self.newprov = {}
        ids = [int(el.get(qn("w:val"))) for el in E.element.body.iter(qn("w:id"))
               if el.getparent() is not None and el.getparent().tag == qn("w:sdtPr")
               and (el.get(qn("w:val")) or "").lstrip("-").isdigit()]
        self.next_sdt_id = (max(ids) + 1) if ids else 200000000

    def owned(self, *keys):
        """What the pipeline last wrote for this cell, under the current key
        or any legacy alias (a column-position or band-name key), or None."""
        for k in keys:
            if k and k.lower() in self.prov_ci:
                return self.prov_ci[k.lower()]
        return None

    def settle(self, keys):
        self.settled.update(k.lower() for k in keys if k)

    def prepare_copy(self, el):
        """Deep-copy a fresh-doc element for insertion into the filed doc:
        re-home any hyperlink relationship and renumber checkbox ids."""
        new = copy.deepcopy(el)
        for link in new.iter(qn("w:hyperlink")):
            rid = link.get(qn("r:id"))
            if not rid:
                continue
            rel = self.F.part.rels.get(rid)
            if rel is None:
                continue
            new_rid = self.E.part.relate_to(rel.target_ref, RT.HYPERLINK, is_external=True)
            link.set(qn("r:id"), new_rid)
        for id_el in new.iter(qn("w:id")):
            if id_el.getparent() is not None and id_el.getparent().tag == qn("w:sdtPr"):
                id_el.set(qn("w:val"), str(self.next_sdt_id))
                self.next_sdt_id += 1
        return new

    def notice(self, section, column, doc_value, new_value, key=None):
        self.notices.append({"field": section + (f" — {column}" if column else ""),
                             "doc_value": doc_value, "new_value": new_value,
                             "cell_key": key})

    def action(self, keys, blank, owned, doc, new, reviewable=False, quiet=False):
        """A cell/paragraph whose doc text differs from the fresh build:
        'write' it, raise a 'notice', or 'skip'. `keys` = the cell's key and
        its legacy aliases (a decision recorded under any of them counts).
        `reviewable` = the cell holds a band answer (review mode doesn't hold
        booking-driven cells like the act-name header or set length — a
        cancellation still lands). `quiet` = never a notice (Consoles, F15)."""
        ks = [k.lower() for k in keys if k]
        notice = "skip" if quiet else "notice"
        fk = next((k for k in ks if k in self.force), None)
        if fk:
            if _t(self.force[fk]) == _t(new):
                self.forced.update(ks)
                return "write"
            return notice     # the answer changed again since Brian looked
        if blank:
            return "write"
        if any(k in self.frozen for k in ks):
            return notice     # waiting on Brian — re-recorded (dedupes; a new value supersedes)
        if owned and not (self.review and reviewable):
            return "write"
        if owned or _differs_meaningfully(doc, new):
            return notice
        return "skip"


def _replace_tc_content(ctx, e_tc, f_tc):
    for child in list(e_tc):
        if child.tag != qn("w:tcPr"):
            e_tc.remove(child)
    for child in f_tc:
        if child.tag != qn("w:tcPr"):
            e_tc.append(ctx.prepare_copy(child))


_TOK = re.compile(r"[a-z0-9]+|[☐☒]")


def _tokens(s):
    return _TOK.findall(_t(s).lower())


def _differs_meaningfully(e_text, f_text):
    """A diff worth reporting: the new value does NOT appear, in order and
    contiguously, inside what the doc says. A hand edit that extends the
    pipeline's value — '5 wedges (2 for drums)' vs '5 wedges' — is not a
    conflict; '3 wedges' vs '5 wedges' is. Whole tokens, not substrings
    (review 2026-09-14: '5' inside '15' used to hide a diff).

    F4 (audit 2026-09-16): this used to be an unordered token SET, so
    '1 large, 2 standard' vs '2 large, 1 standard', a flipped ☐ Yes ☒ No and
    '4 wedges (2 for drums)' vs '2 wedges' all compared equal. Order and
    checkbox glyphs count now."""
    te, tf = _tokens(e_text), _tokens(f_text)
    if not tf:
        return False
    n = len(tf)
    # the new value's last word may run straight into a hand-typed addition
    # ("Split snake: YesBand brings all mics") — still the same value extended
    def same_last(doc_tok):
        return doc_tok == tf[-1] or (tf[-1].isalpha() and doc_tok.startswith(tf[-1])
                                     and doc_tok[len(tf[-1]):].isalpha())
    return not any(te[i:i + n - 1] == tf[:-1] and same_last(te[i + n - 1])
                   for i in range(len(te) - n + 1))


# doc_value of a notice for a pipeline value somebody deleted (F10)
CLEARED_BY_HAND = "(cleared by hand)"
# what "Keep doc" on such a notice stores as the cell's ownership: not the
# template, so the cell keeps reading as deliberately cleared
KEPT_BLANK = "\u2205 kept blank"

# grid rows filled from the booking or house lookups, not the band's form —
# never held for review
BOOKING_ROWS = {"act names", "set length", "consoles"}
# 4-cell rows whose per-column value is a house lookup, not the artist's own:
# a transient staffing-sheet miss must never blank them (no retraction)
HOUSE_ROWS = {"engineer", "consoles"}
# rows that never raise a notice: the pipeline fills what it owns, a person's
# value is left alone silently (F15 — Consoles nagged on every hand value)
QUIET_ROWS = {"consoles"}


def _record_paras(ctx, key, tc):
    """F14 (audit 2026-09-16): alongside the whole-cell key, own each
    paragraph of a multi-paragraph cell — otherwise the first later change to
    one line (a Mon engineer resolving after creation) found no per-paragraph
    ownership and raised a notice instead of filling."""
    paras = _tc_paras(tc)
    if len(paras) > 1 and not _has_table(tc):
        for i, p in enumerate(paras):
            ctx.newprov[f"{key}|{i}"] = _t(_p_text(p))


def _compare_cell(ctx, t_tc, f_tc, e_tc, section, column, key=None, aliases=(),
                  reviewable=False, band_row=False, quiet=False):
    """Merge one cell. `key` is the provenance/notice key — `Section|a<id>`
    for an artist's cell since 2026-09-16, `Section|colN` for the act-name
    header, `Section|` for a one-value row; `aliases` are older spellings of
    the same cell still honoured on read. `band_row` = the value is the
    artist's own (a pipeline value is withdrawn when the fresh build has none)."""
    key = key or f"{section}|"
    keys = [key] + [a for a in aliases if a]
    tT, tF, tE = _t(_tc_text(t_tc)), _t(_tc_text(f_tc)), _t(_tc_text(e_tc))
    prev = ctx.owned(*keys)
    if tF == tT:
        # the pipeline has nothing for this cell
        if prev is not None and prev == tE and tE != tT and band_row \
                and not any(k.lower() in ctx.frozen for k in keys):
            # F7: it was the pipeline's value (a cancelled or removed act, an
            # answer withdrawn) — give the cell back to the template
            _replace_tc_content(ctx, e_tc, t_tc)
            ctx.settle(keys)
            ctx.changed += 1
            return
        if prev == tE:
            ctx.newprov[key] = tE
        return
    if tE == tF:
        ctx.newprov[key] = tF
        _record_paras(ctx, key, e_tc)
        ctx.settle(keys)
        return
    # F10: the pipeline wrote a value and the cell now reads the blank template
    # — a person deleted it on purpose. Ask, don't refill.
    hand_cleared = tE == tT and prev is not None and prev != tT
    cell_owned = (tE == tT and not hand_cleared) or prev == tE
    pT, pF, pE = _tc_paras(t_tc), _tc_paras(f_tc), _tc_paras(e_tc)
    if (not any(_has_table(x) for x in (t_tc, f_tc, e_tc))
            and len(pT) == len(pF) == len(pE) and len(pT) > 1):
        for i, (a, b, c) in enumerate(zip(pT, pF, pE)):
            ta, tb, tc_ = _t(_p_text(a)), _t(_p_text(b)), _t(_p_text(c))
            pkeys = [f"{k}|{i}" for k in keys]
            pk = pkeys[0]
            pprev = ctx.owned(*pkeys)
            if tb == ta:
                if pprev == tc_:
                    ctx.newprov[pk] = tc_
                continue
            if tc_ == tb:
                ctx.newprov[pk] = tb
                ctx.settle(pkeys)
                continue
            p_cleared = tc_ == ta and pprev is not None and pprev != ta and prev != tE
            p_owned = pprev == tc_
            act = ctx.action(pkeys, blank=(tc_ == ta and not p_cleared),
                             owned=((cell_owned and not hand_cleared) or p_owned),
                             doc=tc_, new=tb, reviewable=reviewable, quiet=quiet)
            if act == "write":
                c.addprevious(ctx.prepare_copy(b))
                c.getparent().remove(c)
                ctx.newprov[pk] = tb
                ctx.settle(pkeys)
                ctx.changed += 1
                continue
            if act == "notice":
                ctx.notice(section, column, _p_text(c).strip() or CLEARED_BY_HAND,
                           _p_text(b).strip(), pk)
            if p_owned:
                ctx.newprov[pk] = tc_   # held, not given up: still the pipeline's
            elif p_cleared:
                ctx.newprov[pk] = pprev  # remembered, so the next pass asks again instead of refilling
        if _t(_tc_text(e_tc)) == tF:
            ctx.newprov[key] = tF
        return
    act = ctx.action(keys, blank=(tE == tT and not hand_cleared), owned=cell_owned, doc=tE, new=tF,
                     reviewable=reviewable, quiet=quiet)
    if act == "write":
        _replace_tc_content(ctx, e_tc, f_tc)
        ctx.newprov[key] = tF
        _record_paras(ctx, key, e_tc)
        ctx.settle(keys)
        ctx.changed += 1
        return
    if act == "notice":
        ctx.notice(section, column, _tc_text(e_tc).strip() or CLEARED_BY_HAND,
                   _tc_text(f_tc).strip(), key)
    if prev == tE and not hand_cleared:
        ctx.newprov[key] = tE
    elif hand_cleared:
        # keep the record of what the pipeline wrote: the cell stays "cleared by
        # hand" (one deduped notice) until Brian picks Use new — never a refill
        ctx.newprov[key] = prev


def _rows_by_key(table, label_fn):
    out, counts = {}, {}
    for row in table.rows:
        lab = label_fn(row)
        n = counts.get(lab, 0)
        counts[lab] = n + 1
        out[(lab, n)] = row
    return out


def _grid_label(row):
    tcs = row._tr.tc_lst
    return daysheet.norm(_tc_text(tcs[0])) if tcs else ""


def _act_names(E_grid):
    """column index -> band name as the filed doc shows it (for notices)."""
    names = {}
    for row in E_grid.rows:
        tcs = row._tr.tc_lst
        if not tcs or _t(_tc_text(tcs[0])):
            continue
        for i, tc in enumerate(tcs[1:], start=1):
            paras = _tc_paras(tc)
            if len(paras) >= 2 and _t(_p_text(paras[1])):
                names[i] = _t(_p_text(paras[1]))
        if names:
            break
    return names


def _int_keys(m):
    out = {}
    for k, v in (m or {}).items():
        try:
            if v is not None:
                out[int(k)] = int(v)
        except (TypeError, ValueError):
            continue
    return out


def _rehome_columns(ctx, gT, gE, old_cols, new_cols):
    """Physically move each artist's cells to the column they occupy now
    (root cause B, audit 2026-09-16). For every 4-cell row of the grid except
    the act-name header (whose ARTIST 1/2/3 labels are positional), each
    column's content — everything but its w:tcPr, hand edits included — is
    lifted and dropped into the column its artist moved to. A column whose
    artist left and that nobody moved into gets the template cell back.
    Returns the number of rows touched."""
    old_col_of = {aid: c for c, aid in old_cols.items()}
    plan = {}
    for i in (1, 2, 3):
        aid = new_cols.get(i)
        if aid is not None and aid in old_col_of:
            if old_col_of[aid] != i:
                plan[i] = old_col_of[aid]
        elif old_cols.get(i) is not None:
            plan[i] = None
    if not plan:
        return 0
    t_rows = _rows_by_key(gT, _grid_label)
    touched = 0
    for (lab, n), e_row in _rows_by_key(gE, _grid_label).items():
        tcs = e_row._tr.tc_lst
        if len(tcs) != 4 or lab == "":
            continue
        t_row = t_rows.get((lab, n))
        t_tcs = t_row._tr.tc_lst if t_row is not None else []
        saved = {i: [copy.deepcopy(ch) for ch in tcs[i] if ch.tag != qn("w:tcPr")] for i in (1, 2, 3)}
        for i, src in plan.items():
            for ch in list(tcs[i]):
                if ch.tag != qn("w:tcPr"):
                    tcs[i].remove(ch)
            if src is not None:
                for ch in saved[src]:
                    tcs[i].append(ch)
            elif len(t_tcs) == 4:
                for ch in t_tcs[i]:
                    if ch.tag != qn("w:tcPr"):
                        tcs[i].append(ctx.prepare_copy(ch))
            if not any(ch.tag == qn("w:p") for ch in tcs[i]):
                tcs[i].append(OxmlElement("w:p"))
        touched += 1
    return touched


def header_column_map(doc, acts=None):
    """{column: artist_id} read off a filed doc's own act-name header — for a
    doc registered before filed_docs.columns existed. Names match the bill's
    acts first, then any artist on file by match_key; a column whose name
    matches nobody is left out."""
    grid = daysheet.find_grid(doc)
    names = _act_names(grid) if grid is not None else {}
    by_name = {db.normalize(a["artist"]["name"]): a["artist_id"]
               for a in acts or [] if a.get("artist") and a.get("artist_id")}
    out, missing = {}, {}
    for i, nm in names.items():
        n = db.normalize(re.sub(r"^cancelled\s*[—–-]\s*", "", nm, flags=re.I))
        if n in by_name:
            out[i] = by_name[n]
        elif n:
            missing[i] = n
    if missing:
        with db.get_conn() as conn, conn.cursor() as cur:
            cur.execute("SELECT id, match_key FROM artists WHERE match_key = ANY(%s)",
                        (list(missing.values()),))
            ids = {r["match_key"]: r["id"] for r in cur.fetchall()}
        out.update({i: ids[n] for i, n in missing.items() if n in ids})
    return out


def merge(T, F, E, prov=None, review=False, frozen=None, force=None, info=None,
          cols=None, old_cols=None):
    """Update E (filed doc) from F (fresh build): blank cells and cells the
    pipeline itself last wrote (per `prov`) take the fresh value; anything a
    person typed is left alone and reported. review=True holds pipeline-owned
    cells too; `frozen` keys have an open notice; `force` {key: new text} are
    Brian's "Use new" picks. `cols` is the fresh build's {column: artist_id}
    and `old_cols` what the filed doc was last written with — when they
    differ, E's cells are re-homed first. `info` (a dict) receives the
    settled/forced key sets. Returns (changes_made, notices, new_provenance)."""
    ctx = _Ctx(T, F, E, prov=prov, review=review, frozen=frozen, force=force)
    cols, old_cols = _int_keys(cols), _int_keys(old_cols)

    gT, gF, gE = daysheet.find_grid(T), daysheet.find_grid(F), daysheet.find_grid(E)
    if gT is not None and gF is not None and gE is not None:
        if cols and old_cols and cols != old_cols and _rehome_columns(ctx, gT, gE, old_cols, cols):
            ctx.changed += 1
        rT, rF, rE = (_rows_by_key(g, _grid_label) for g in (gT, gF, gE))
        # F13: label notices from the FRESH header — the filed one still names
        # whoever used to sit in the column on the pass that reorders the bill
        names = _act_names(gF)
        old_col_of = {aid: c for c, aid in old_cols.items()}
        for key, t_row in rT.items():
            f_row, e_row = rF.get(key), rE.get(key)
            if f_row is not None and e_row is None and key[1] == 0:
                # a row the template gained after this doc was filed (e.g.
                # Additional Info, 2026-09-14) — append it from the fresh build
                gE._tbl.append(ctx.prepare_copy(f_row._tr))
                for i, tc in enumerate(f_row._tr.tc_lst[1:], start=1):
                    v = _t(_tc_text(tc))
                    if v:
                        sec = _t(_tc_text(f_row._tr.tc_lst[0])).rstrip(':')
                        aid = cols.get(i)
                        ctx.newprov[f"{sec}|a{aid}" if aid is not None else f"{sec}|col{i}"] = v
                ctx.changed += 1
                continue
            if f_row is None or e_row is None:
                continue
            tT, tF, tE = t_row._tr.tc_lst, f_row._tr.tc_lst, e_row._tr.tc_lst
            if not (len(tT) == len(tF) == len(tE)):
                continue
            section = _t(_tc_text(tT[0])).rstrip(":") or "Act names"
            low = section.lower()
            if len(tT) == 2:
                _compare_cell(ctx, tT[1], tF[1], tE[1], section, None,
                              reviewable=low not in BOOKING_ROWS)
                continue
            header = key[0] == ""
            for i in range(1, len(tT)):
                col = names.get(i) or ACT_COL_NAMES.get(i)
                aid = cols.get(i)
                if header or aid is None:
                    ckey, aliases = f"{section}|col{i}", [f"{section}|{col}"]
                else:
                    ckey = f"{section}|a{aid}"
                    aliases = [f"{section}|{col}"]
                    # pre-2026-09-16 position key: where this artist sat when it
                    # was recorded. With no stored map the column can't have
                    # been tracked, so its own position is the best guess.
                    if aid in old_col_of:
                        aliases.append(f"{section}|col{old_col_of[aid]}")
                    elif not old_cols:
                        aliases.append(f"{section}|col{i}")
                _compare_cell(ctx, tT[i], tF[i], tE[i], section, col, key=ckey, aliases=aliases,
                              reviewable=low not in BOOKING_ROWS,
                              # a column nobody occupies any more, the positional
                              # header, and every artist-owned row give a stale
                              # pipeline value back; a house lookup under a real
                              # artist never does (a staffing-sheet miss)
                              band_row=header or aid is None or low not in HOUSE_ROWS,
                              quiet=low in QUIET_ROWS)

    lT, lF, lE = daysheet.find_lead_table(T), daysheet.find_lead_table(F), daysheet.find_lead_table(E)
    if lT is not None and lF is not None and lE is not None:
        for ri in range(min(len(lT.rows), len(lF.rows), len(lE.rows))):
            a, b, c = lT.rows[ri]._tr.tc_lst, lF.rows[ri]._tr.tc_lst, lE.rows[ri]._tr.tc_lst
            if len(a) >= 2 and len(b) >= 2 and len(c) >= 2:
                _compare_cell(ctx, a[1], b[1], c[1], _t(_tc_text(a[0])) or "Lead", None)

    sT, sF, sE = (daysheet.find_schedule_table(d) for d in (T, F, E))
    if sT is not None and sF is not None and sE is not None:
        # 2026-09-16 rewrite. This used to decide what to do by comparing the
        # schedule's LABEL LISTS (template vs fresh vs filed) and did nothing at
        # all unless they lined up. That held while labels were fixed template
        # text. Since 2026-09-15 the labels carry artist names ("Joe Jordan
        # Load-In") and drop the half of a row nobody plays — so they change the
        # moment the bill changes, and adding a second artist to a filed doc
        # skipped the whole schedule: Augustana (FSQ 9/25) reached the AUDIO
        # columns and never reached the day's schedule.
        #
        # Now the table is matched by SHAPE. Same row count means row i is the
        # same slot of the day in both docs: its label is pipeline-owned and
        # always rewritten from the fresh build (nobody hand-types "Joe Jordan
        # Load-In"), and its time goes through the ordinary cell merge, keyed by
        # row position so a rename can't orphan it.
        def label_of(row):
            tcs = row._tr.tc_lst
            return _t(_tc_text(tcs[-1])) if tcs else ""

        def time_of(row):
            tcs = row._tr.tc_lst
            return _t(_tc_text(tcs[0])) if tcs else ""

        def time_owned(i, row, old_label):
            v = time_of(row)
            return bool(v) and (ctx.prov.get(f"Schedule|row{i}") == v
                                or ctx.owned(f"Schedule row{i}|", f"Schedule: {old_label}|") == v)

        if len(sE.rows) == len(sF.rows):
            for i, (rf, re_) in enumerate(zip(sF.rows, sE.rows)):
                f_tcs, e_tcs = rf._tr.tc_lst, re_._tr.tc_lst
                if not f_tcs or not e_tcs:
                    continue
                old_label = label_of(re_)
                if label_of(rf) != old_label:
                    _replace_tc_content(ctx, e_tcs[-1], f_tcs[-1])
                    ctx.changed += 1
                if not time_of(rf) and time_owned(i, re_, old_label):
                    # the artist this time belonged to is no longer on the bill
                    # and the pipeline wrote it — clear it rather than leave a
                    # time next to a row that now names nobody
                    _replace_tc_content(ctx, e_tcs[0], f_tcs[0])
                    ctx.changed += 1
                    continue
                _compare_cell(ctx, _blank_like(f_tcs[0]), f_tcs[0], e_tcs[0],
                              f"Schedule row{i}", None,
                              aliases=[f"Schedule: {old_label}|"])
                v = time_of(re_)
                if v and ctx.prov.get(f"Schedule|row{i}") == v:
                    ctx.newprov[f"Schedule|row{i}"] = v
        elif all(not time_of(r) or time_owned(i, r, label_of(r))
                 for i, r in enumerate(sE.rows)):
            # Different shape — a series schedule was locked or unlocked since
            # this doc was filed — and nothing in the filed table is a human's:
            # replace the rows wholesale.
            tbl_el = sE._tbl
            for r in list(sE.rows):
                tbl_el.remove(r._tr)
            for i, r in enumerate(sF.rows):
                tbl_el.append(ctx.prepare_copy(r._tr))
                ctx.newprov[f"Schedule|row{i}"] = _t(_tc_text(r._tr.tc_lst[0]))
            ctx.changed += 1
        else:
            # Different shape AND hand-typed times in the filed table. Silently
            # skipping this is how a schedule goes stale unnoticed, so say so.
            # F11: the fresh row labels ride in new_value, so a schedule whose
            # shape changes AGAIN after a kept notice is a new notice, not a
            # duplicate the UNIQUE index swallows forever
            ctx.notice("Schedule", None,
                       "hand-typed times in a schedule whose rows no longer match",
                       "rebuild the schedule block by hand, or clear it and re-run — new rows: "
                       + " · ".join(label_of(r) for r in sF.rows),
                       "Schedule|shape")
    if info is not None:
        info["settled"], info["forced"] = ctx.settled, ctx.forced
    return ctx.changed, ctx.notices, ctx.newprov


def all_acts_cancelled(ev, acts):
    """Every act on this bill has its show cancelled (review 2026-09-14, M4)
    -> the filed doc is renamed with a CANCELLED marker. One cancelled band
    on a multi-band bill only blanks its own column (daysheet.build)."""
    ids = [a.get("artist_id") for a in acts if a.get("artist_id")]
    if not ids or not ev.get("event_date"):
        return False
    with db.get_conn() as conn, conn.cursor() as cur:
        cur.execute("""SELECT artist_id FROM shows WHERE venue=%s AND show_date=%s
                       AND cancelled_at IS NOT NULL AND artist_id = ANY(%s)""",
                    (ev.get("venue"), ev.get("event_date"), ids))
        cancelled = {r["artist_id"] for r in cur.fetchall()}
    return all(i in cancelled for i in ids)


def _blank_like(tc):
    """A stand-in 'template' cell with no text, for locked-schedule rows that
    have no pristine template equivalent."""
    new = copy.deepcopy(tc)
    for t in new.iter(qn("w:t")):
        t.text = ""
    return new


# ── filing ──────────────────────────────────────────────────────────────────

def _atomic_save(doc, dest, expect_sha=None):
    """Save `doc` to `dest` via a temp file in the same folder. If
    expect_sha is given and the file on disk no longer matches it (someone
    saved it while we were working), abort — never clobber."""
    dest = Path(dest)
    tmp = dest.parent / f".~advtmp-{os.getpid()}-{dest.name}"
    doc.save(str(tmp))
    if expect_sha is not None and dest.exists() and sha256_file(dest) != expect_sha:
        tmp.unlink(missing_ok=True)
        return False
    if word_lock_present(dest):
        tmp.unlink(missing_ok=True)
        return False
    os.replace(tmp, dest)
    return True


def _rel_to_root(path):
    try:
        return str(Path(path).relative_to(fs.real_dropbox_root()))
    except ValueError:
        return str(path)


def _abs_from_reg(p):
    p = Path(p)
    return p if p.is_absolute() else fs.real_dropbox_root() / p


def locate(cur, venue, event_date, key, n_events_same_day=1):
    """The registry row for this show's doc, or None. Exact key first; if
    the event name was edited (key changed) and this venue+date has only
    one registered doc and one event, that doc is still this show's."""
    row = db.get_filed_doc(cur, venue, event_date, key)
    if row:
        return row
    rows = db.filed_docs_for(cur, venue, event_date)
    if len(rows) == 1 and n_events_same_day == 1:
        return rows[0]
    return None


def _find_moved_doc(folder, d, reg):
    """The doc a registry row points at, after someone renamed it: first a
    file with the exact recorded hash, else the single date-stamped doc in
    the folder that no other registry row claims."""
    if not folder.exists():
        return None
    cands = sorted(folder.glob(f"{d.strftime('%m%d%y')} * Prod Adv.docx"))
    if reg.get("sha256"):
        for c in cands:
            if sha256_file(c) == reg["sha256"]:
                return c
    with db.get_conn() as conn, conn.cursor() as cur:
        claimed = {(_abs_from_reg(r["path"])).name for r in db.filed_docs_for(cur, reg["venue"], d)
                   if r["id"] != reg["id"]}
    free = [c for c in cands if c.name not in claimed]
    return free[0] if len(free) == 1 else None


def file_event_doc(eid, ev, acts, stem, stageplot_names=None, today=None, dry_run=False,
                   n_events_same_day=1, force=None):
    """Create or merge-fill this event's filed doc. Returns a result dict:
    action in created|merged|unchanged|locked|past|conflict|dry-run, path,
    notices (list), hand_edited (bool), renamed_from (str|None), review
    (bool), applied (keys written from `force`)."""
    today = today or dt.date.today()
    d = ev.get("event_date")
    res = {"action": None, "path": None, "notices": [], "hand_edited": False,
           "renamed_from": None, "filled": 0, "review": False, "applied": set()}
    if not d:
        res["action"] = "no-date"
        return res
    if d < today:
        res["action"] = "past"
        return res
    venue = ev.get("venue")
    folder = fs.real_dropbox_root() / fs.real_venue_folder(venue) / fs.real_month_folder(venue, d)
    desired = folder / f"{stem}.docx"
    key = event_key(ev)

    with db.get_conn() as conn, conn.cursor() as cur:
        reg = locate(cur, venue, d, key, n_events_same_day)
        open_notices = db.open_doc_notices(cur, venue, d, key) if reg else []
        # F11: a doc whose last decision was made after its last write is
        # reviewed as of that decision, not stuck in review mode for good
        since = max([t for t in ((reg or {}).get("written_at"), (reg or {}).get("reviewed_at")) if t],
                    default=None)
        review = bool(reg) and db.resubmitted_since(
            cur, venue, d, [a.get("artist_id") for a in acts if a.get("artist_id")], since)
    res["review"] = review
    existing = None
    if reg and _abs_from_reg(reg["path"]).exists():
        existing = _abs_from_reg(reg["path"])
    elif ci_existing(folder, desired.name):
        existing = ci_existing(folder, desired.name)
    elif reg:
        # registered doc is gone from where we left it — most likely a person
        # renamed or moved it within the month folder. Adopt the one
        # unregistered doc for this date rather than creating a duplicate.
        existing = _find_moved_doc(folder, d, reg)
    res["path"] = str(existing or desired)

    fresh, _event, _acts, _n = daysheet.build(eid, stageplot_names=stageplot_names or {})
    cols = daysheet.column_map(_acts, _n)
    extra = daysheet.acts_off_the_doc(_acts, _n)
    if extra:
        # F9: never folded onto column 3 — say so instead
        res["notices"].append({
            "field": "Bill size", "kind": "bill_size",
            "doc_value": "template holds 3 artists",
            "new_value": f"{len(_acts)} acts, template holds 3 — not on the doc: "
                         + ", ".join(a["artist"]["name"] for a in extra),
            "cell_key": f"bill_size:{len(_acts)}"})

    if existing is None:
        res["action"] = "created"
        if dry_run:
            return res
        folder.mkdir(parents=True, exist_ok=True)
        if not _atomic_save(fresh, desired):
            res["action"] = "locked"
            return res
        # provenance of a freshly created doc: every cell that differs from
        # the template is the pipeline's (merge of the fresh doc against
        # itself records exactly those)
        _c, _nn, prov = merge(Document(str(daysheet.UNIVERSAL_TEMPLATE)),
                              Document(str(desired)), Document(str(desired)), cols=cols)
        with db.get_conn() as conn, conn.cursor() as cur:
            db.upsert_filed_doc(cur, venue, d, key, _rel_to_root(desired), sha256_file(desired),
                                cells=prov, columns=cols)
            conn.commit()
        res["path"] = str(desired)
        return res

    for cc in conflicted_copies(existing):
        # F2: somebody saved this doc in Word while the pipeline wrote it — their
        # edits are only in the copy. Filing carries on (blank fills lose nothing
        # more); Brian merges the copy back by hand.
        res["notices"].append({"field": "Conflicted copy", "kind": "conflict",
                               "doc_value": existing.name, "new_value": cc.name,
                               "cell_key": f"conflict:{cc.name}"})
    if word_lock_present(existing):
        res["action"] = "locked"
        return res

    before_sha = sha256_file(existing)
    res["hand_edited"] = bool(reg and reg.get("sha256") and reg["sha256"] != before_sha)
    E = Document(str(existing))
    T = Document(str(daysheet.UNIVERSAL_TEMPLATE))
    info = {}
    changed, notices, newprov = merge(T, fresh, E, prov=(reg or {}).get("cells") or {},
                                      review=review, force=force, info=info,
                                      frozen=[n["cell_key"] for n in open_notices if n.get("cell_key")],
                                      cols=cols,
                                      old_cols=(reg or {}).get("columns") or header_column_map(E, _acts))
    res["notices"] = res["notices"] + notices
    res["applied"] = info["forced"]
    res["filled"] = changed
    res["action"] = "merged" if changed else "unchanged"
    if dry_run:
        return res

    target = existing
    if changed:
        if not _atomic_save(E, existing, expect_sha=before_sha):
            res["action"] = "conflict"
            return res
    # rename in place when the display name changed — never a second copy,
    # and never a case-only rename (Dropbox on the Mac can't tell them apart)
    if (existing.name != desired.name and existing.name.lower() != desired.name.lower()
            and not ci_existing(folder, desired.name)):
        folder.mkdir(parents=True, exist_ok=True)
        os.rename(existing, desired)
        res["renamed_from"] = existing.name
        target = desired
    res["path"] = str(target)
    with db.get_conn() as conn, conn.cursor() as cur:
        new_sha = sha256_file(target)
        # keep the recorded hash as "what the pipeline last wrote" — only
        # refresh it when the pipeline itself wrote (or first registers)
        if changed or not reg:
            db.upsert_filed_doc(cur, venue, d, key, _rel_to_root(target), new_sha,
                                seeded=not reg and not changed, reg_id=(reg or {}).get("id"),
                                cells=newprov, columns=cols)
        else:
            db.upsert_filed_doc(cur, venue, d, key, _rel_to_root(target), reg.get("sha256"),
                                reg_id=reg.get("id"), keep_written_at=True, cells=newprov,
                                columns=cols)
        # an open decision whose cell now reads the fresh value is done:
        # applied if Brian chose it, otherwise moot (the doc and the form agree)
        for n in open_notices:
            k = (n.get("cell_key") or "").lower()
            if k and k in info["settled"]:
                db.resolve_doc_notice(cur, n["id"], "applied" if k in info["forced"] else "moot")
        conn.commit()
    return res


def file_stage_plot(src, folder, fname, dry_run=False):
    """Copy a band's uploaded plot next to the doc. Never overwrites a
    different existing plot: a changed upload is saved alongside as
    '<stem> (updated MMDDYY-HHMM)<ext>' and reported. Returns
    (filed_name, notice_or_None)."""
    folder = Path(folder)
    dest = folder / fname
    if dry_run:
        return fname, None
    folder.mkdir(parents=True, exist_ok=True)
    existing = ci_existing(folder, fname)
    if existing is not None:
        dest, fname = existing, existing.name   # keep whatever case is already there
    if not dest.exists():
        shutil.copy(src, dest)
        return fname, None
    if filecmp.cmp(src, dest, shallow=False):
        return fname, None
    stem, ext = Path(fname).stem, Path(fname).suffix
    for p in folder.glob(f"{stem} (updated *){ext}"):
        if filecmp.cmp(src, p, shallow=False):
            return fname, None
    alt = f"{stem} (updated {dt.datetime.now().strftime('%m%d%y-%H%M')}){ext}"
    shutil.copy(src, folder / alt)
    # kind="stage_plot" (audit 2026-09-16 #23): this is an FYI, not a
    # reviewable conflict — there's no doc cell to "Keep" or "Apply" for a
    # separate uploaded file — so it's auto-resolved at insert instead of
    # sitting in the kind='diff' review queue forever with nothing that
    # would ever resolve it.
    return fname, {"field": "Stage plot file", "doc_value": fname,
                   "new_value": f"new upload saved as {alt}", "cell_key": f"plot:{alt}",
                   "kind": "stage_plot"}


def record_and_email_notices(results, send_mail=True):
    """results: [(ev, res)]. Stores each notice once (doc_notices), then
    queues ONE digest item with the notices not reported before (review
    2026-09-14, E1: these ride the 7am digest and the dashboard's "Needs
    you" panel, not their own email). Returns the number of new notices."""
    import mailer
    inserted = 0
    with db.get_conn() as conn, conn.cursor() as cur:
        for ev, res in results:
            for n in res.get("notices") or []:
                kind = n.get("kind") or "diff"
                nid = db.insert_doc_notice(cur, ev.get("venue"), ev.get("event_date"),
                                           event_key(ev), kind, n["field"],
                                           n.get("new_value") or "", n.get("doc_value") or "",
                                           detail=Path(res.get("path") or "").name
                                           + (" (hand-edited)" if res.get("hand_edited") else ""),
                                           cell_key=n.get("cell_key"))
                if nid:
                    inserted += 1
                    # audit 2026-09-16 #23: only kind='diff' is a real
                    # decision for Brian (open_doc_notices/the review UI
                    # both filter on it) — anything else is informational
                    # and would otherwise sit "open" forever with no path
                    # that ever resolves it.
                    if kind != "diff":
                        db.resolve_doc_notice(cur, nid, "auto")
        conn.commit()
        pending = db.unnotified_doc_notices(cur)
    if not pending or not send_mail:
        return inserted
    new = pending
    rows = []
    for n in new:
        where = (f"{mailer.esc(n['venue'])} {n['event_date'].strftime('%m/%d')} — "
                 f"{mailer.esc(n.get('detail'))}")
        edited = ""
        rows.append(
            f"<tr><td style='padding:6px 10px;vertical-align:top'>{where}{edited}</td>"
            f"<td style='padding:6px 10px;vertical-align:top'><b>{mailer.esc(n['field'])}</b></td>"
            f"<td style='padding:6px 10px;vertical-align:top;white-space:pre-wrap'>{mailer.esc(n.get('doc_value'))}</td>"
            f"<td style='padding:6px 10px;vertical-align:top;white-space:pre-wrap'>{mailer.esc(n.get('new_value'))}</td></tr>")
    public = os.environ.get("ADVANCE_PUBLIC_URL", "https://advance.tinydoorstudios.com")
    shows = {(n["venue"], n["event_date"]) for n in new if n.get("event_date")}
    links = "".join(
        f"<li><a href='{public}/doc-review?venue={quote(v)}&date={d.isoformat()}'>"
        f"{mailer.esc(v)} {d.strftime('%m/%d')}</a></li>" for v, d in sorted(shows, key=lambda x: (x[1], x[0])))
    html = ("<div style='font-family:-apple-system,sans-serif;font-size:13px'>"
            "<p>These filed advance docs were <b>not changed</b> — the doc already has a value, "
            "and newer info came in that's different. Each one is waiting on your call "
            "(Keep doc / Use new):</p><ul>" + links + "</ul>"
            "<table style='border-collapse:collapse' border='1' cellspacing='0'>"
            "<tr><th>Show / doc</th><th>Field</th><th>Doc says</th><th>New info</th></tr>"
            + "".join(rows) + "</table></div>")
    with db.get_conn() as conn, conn.cursor() as cur:
        db.queue_digest_item(cur, "doc", f"Advance doc changes to review — {len(new)} field(s)", html)
        db.mark_doc_notices_notified(cur, [n["id"] for n in new])
        conn.commit()
    return inserted
