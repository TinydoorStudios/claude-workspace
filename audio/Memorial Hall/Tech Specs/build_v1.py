"""Memorial Hall Technical Specifications, Version 1 (2026-09-29).

Source: 'Memorial Hall Production Specifications - 2026 edit.docx' (Brian's Downloads).
Brand: Code/BandInfoForm/app/brand/brand.json (Avenir Next, Memo gray + purple, logos).
Builds the .docx here; build_pdf() has Word export the PDF so both files match.

    python3 build_v1.py
"""
import json
import os
import subprocess

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
BRAND_DIR = os.path.expanduser("~/Documents/Claude/Code/BandInfoForm/app/brand")
BRAND = json.load(open(os.path.join(BRAND_DIR, "brand.json")))
VENUE = BRAND["venues"]["Memorial Hall"]

VERSION = "Version 1"
DATE = "September 2026"
OUT = os.path.join(HERE, "Memorial Hall Technical Specifications v1")

BODY = BRAND["fonts"]["body"]
HEAD = BRAND["fonts"]["heading"]
PURPLE = VENUE["secondary"]["hex"].lstrip("#")
GRAY = VENUE["primary"]["hex"].lstrip("#")
TEXT = BRAND["corporate"]["text"].lstrip("#")
TINT = "F1EEF3"      # label cells: a whisper of the Memo purple
RULE = "D9D9DB"

CONTENT_W = 7.3      # inches, letter with 0.6" margins


def rgb(h):
    return RGBColor.from_string(h)


def photo(n):
    return os.path.join(HERE, "photos", f"image{n}.jpg")


# ---------------------------------------------------------------- low-level XML

def set_run_font(run, name=BODY, size=None, bold=None, italic=None, color=TEXT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = rgb(color)
    return run


def shade(cell, hexcolor):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)


def table_borders(table, color=RULE, size=4, inside_v=False, outer=True):
    tblpr = table._element.tblPr
    b = OxmlElement("w:tblBorders")
    edges = ["top", "left", "bottom", "right", "insideH", "insideV"]
    for e in edges:
        el = OxmlElement(f"w:{e}")
        on = (e == "insideH") or (e == "insideV" and inside_v) or (e in ("top", "bottom") and outer)
        el.set(qn("w:val"), "single" if on else "nil")
        if on:
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), color)
            el.set(qn("w:space"), "0")
        b.append(el)
    tblpr.append(b)


def no_borders(table):
    tblpr = table._element.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "nil")
        b.append(el)
    tblpr.append(b)


def cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblpr = table._element.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tblpr.append(m)


def fixed_widths(table, widths):
    table.autofit = False
    tblpr = table._element.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w)
    grid = table._element.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trpr.append(el)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    trpr.append(el)


def para_border_bottom(p, color=PURPLE, size=12, space=2):
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), str(size))
    bot.set(qn("w:space"), str(space))
    bot.set(qn("w:color"), color)
    pbdr.append(bot)
    ppr.append(pbdr)


def keep_next(p):
    p.paragraph_format.keep_with_next = True


def spacing(p, before=0, after=0, line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line


def add_field(run, instr):
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            run._r.append(f)
        elif text == instr:
            t = OxmlElement("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = f" {instr} "
            run._r.append(t)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)


def hyperlink(p, url, text, size=9, color=PURPLE):
    part = p.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), BODY)
    rpr.append(rf)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rpr.append(c)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    p._p.append(h)


# ---------------------------------------------------------------- document setup

def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.6))
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.6)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.3)

    # docDefaults + Normal so nothing inherits Calibri/Times
    styles = doc.styles
    rpr_default = styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    rf = rpr_default.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr_default.insert(0, rf)
    for a in list(rf.attrib):
        del rf.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), BODY)
    normal = styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = rgb(TEXT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.1
    for s in ("Heading 1", "Heading 2"):
        st = styles[s]
        st.font.name = HEAD
        st.font.color.rgb = rgb(PURPLE if s == "Heading 1" else TEXT)
        st.font.size = Pt(17 if s == "Heading 1" else 11)
        st.font.bold = True
        st.font.italic = False
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for a in list(rf.attrib):
            del rf.attrib[a]
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), HEAD)

    # header: Memo wordmark left, doc title right, purple rule under
    hdr = sec.header
    hp = hdr.paragraphs[0]
    t = hdr.add_table(rows=1, cols=2, width=Inches(CONTENT_W))
    no_borders(t)
    fixed_widths(t, [3.6, CONTENT_W - 3.6])
    lp = t.cell(0, 0).paragraphs[0]
    lp.add_run().add_picture(os.path.join(BRAND_DIR, "logos", "memorial-hall-typelogo.png"), height=Inches(0.3))
    rp = t.cell(0, 1).paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(rp.add_run("TECHNICAL SPECIFICATIONS"), HEAD, 10.5, True, color=TEXT)
    rp2 = t.cell(0, 1).add_paragraph()
    rp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(rp2.add_run(f"{VERSION}  ·  {DATE}"), BODY, 8, color=GRAY)
    for c in t.rows[0].cells:
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    # move the empty default paragraph below the table and use it as the rule
    hdr._element.remove(hp._p)
    rule = hdr.add_paragraph()
    spacing(rule, 0, 0)
    para_border_bottom(rule, PURPLE, 12, 1)
    rule.paragraph_format.line_spacing = Pt(4)

    # footer: 3CDC mark left, address + page right
    ftr = sec.footer
    fp = ftr.paragraphs[0]
    t = ftr.add_table(rows=1, cols=2, width=Inches(CONTENT_W))
    no_borders(t)
    fixed_widths(t, [1.6, CONTENT_W - 1.6])
    t.cell(0, 0).paragraphs[0].add_run().add_picture(
        os.path.join(BRAND_DIR, "logos", "3cdc-secondary-horizontal.png"), height=Inches(0.26))
    rp = t.cell(0, 1).paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(rp.add_run(f"Memorial Hall  ·  {VENUE['address']}  ·  memorialhallotr.com     Page "),
                 BODY, 7.5, color=GRAY)
    r = set_run_font(rp.add_run(), BODY, 7.5, color=GRAY)
    add_field(r, "PAGE")
    set_run_font(rp.add_run(" of "), BODY, 7.5, color=GRAY)
    r = set_run_font(rp.add_run(), BODY, 7.5, color=GRAY)
    add_field(r, "NUMPAGES")
    for c in t.rows[0].cells:
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    ftr._element.remove(fp._p)


# ---------------------------------------------------------------- building blocks

def h1(doc, text, page_break=False):
    p = doc.add_paragraph(style="Heading 1")
    if page_break:
        p.paragraph_format.page_break_before = True
    spacing(p, 14, 6)
    set_run_font(p.add_run(text.upper()), HEAD, 17, True, color=PURPLE)
    para_border_bottom(p, GRAY, 6, 3)
    keep_next(p)
    return p


def h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    spacing(p, 9, 3)
    set_run_font(p.add_run(text.upper()), HEAD, 11, True, color=TEXT)
    keep_next(p)
    return p


def body(doc, text, size=9.5, after=4, italic=False, color=TEXT):
    p = doc.add_paragraph()
    spacing(p, 0, after)
    set_run_font(p.add_run(text), BODY, size, italic=italic, color=color)
    return p


def rich(cell_or_doc, parts, size=9.5, first=True):
    """parts: list of (text, bold) tuples."""
    p = cell_or_doc.paragraphs[0] if first else cell_or_doc.add_paragraph()
    for text, b in parts:
        set_run_font(p.add_run(text), BODY, size, bold=b)
    return p


def spec_table(doc, rows, label_w=1.65):
    """Two-column label/value table. value may be a str or a list of lines."""
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t)
    cell_margins(t)
    for label, value in rows:
        r = t.add_row()
        no_split(r)
        lc, vc = r.cells
        shade(lc, TINT)
        set_run_font(lc.paragraphs[0].add_run(label), BODY, 9, True, color=TEXT)
        lines = value if isinstance(value, list) else [value]
        for i, line in enumerate(lines):
            p = vc.paragraphs[0] if i == 0 else vc.add_paragraph()
            if i:
                spacing(p, 2, 0)
            set_run_font(p.add_run(line), BODY, 9.5)
    fixed_widths(t, [label_w, CONTENT_W - label_w])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def grid_table(doc, header, rows, widths, size=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t)
    cell_margins(t, 50, 50, 100, 100)
    hr = t.rows[0]
    repeat_header(hr)
    for c, text in zip(hr.cells, header):
        shade(c, PURPLE)
        set_run_font(c.paragraphs[0].add_run(text.upper()), HEAD, 9, True, color="FFFFFF")
    for row in rows:
        r = t.add_row()
        no_split(r)
        for c, text in zip(r.cells, row):
            set_run_font(c.paragraphs[0].add_run(str(text)), BODY, size)
    fixed_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def photos(doc, items, height=None):
    """items: list of (image number, caption). One row, equal columns."""
    n = len(items)
    gap = 0.15
    col = (CONTENT_W - gap * (n - 1)) / n
    t = doc.add_table(rows=1, cols=n)
    no_borders(t)
    cell_margins(t, 0, 0, 0, 0)
    widths = [col] * n
    fixed_widths(t, widths)
    no_split(t.rows[0])
    for i, (num, cap) in enumerate(items):
        c = t.cell(0, i)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        w = col - (0.08 if n > 1 else 0)
        if height:
            p.add_run().add_picture(photo(num), height=Inches(height))
        else:
            p.add_run().add_picture(photo(num), width=Inches(w))
        cp = c.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing(cp, 3, 0)
        set_run_font(cp.add_run(cap), BODY, 7.5, italic=True, color=GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


# ---------------------------------------------------------------- content

def page_one(doc):
    p = doc.add_paragraph()
    spacing(p, 4, 0)
    set_run_font(p.add_run("MEMORIAL HALL"), HEAD, 30, True, color=TEXT)
    p = doc.add_paragraph()
    spacing(p, 0, 4)
    set_run_font(p.add_run("Technical Specifications"), HEAD, 15, False, color=PURPLE)
    p = doc.add_paragraph()
    spacing(p, 0, 10)
    set_run_font(p.add_run(f"{VENUE['address']}   ·   "), BODY, 9)
    hyperlink(p, "https://www.memorialhallotr.com", "memorialhallotr.com")
    set_run_font(p.add_run("   ·   Virtual tour: "), BODY, 9)
    hyperlink(p, "https://www.memorialhallotr.com/virtual-tour/", "memorialhallotr.com/virtual-tour")

    # contacts
    t = doc.add_table(rows=1, cols=2)
    no_borders(t)
    cell_margins(t, 90, 90, 140, 140)
    contacts = [("Director of Memorial Hall", "Joshua Steele", "JSteele@3cdc.org"),
                ("Senior Production Supervisor", "Brian Lloyd", "BLloyd@3cdc.org")]
    for c, (role, name, email) in zip(t.rows[0].cells, contacts):
        shade(c, TINT)
        p = c.paragraphs[0]
        set_run_font(p.add_run(role.upper()), HEAD, 8.5, True, color=PURPLE)
        p = c.add_paragraph()
        set_run_font(p.add_run(name), BODY, 11, True)
        p = c.add_paragraph()
        hyperlink(p, f"mailto:{email}", email, 9, TEXT)
    fixed_widths(t, [CONTENT_W / 2 - 0.05, CONTENT_W / 2 - 0.05])
    body(doc, "", after=2)

    h1(doc, "At a Glance")
    spec_table(doc, [
        ("Capacity", "556 seats: 215 on the main level (orchestra) and 341 in the balcony. All 11 accessible seats are "
                     "on the main level."),
        ("Stage", "Trapezoidal, 37'4\" at the widest point and 22'5\" deep at center. Permanent orchestra shell; "
                  "no soft goods and no overstage rigging."),
        ("Load-in", "Street-level door on Grant Street (2'7\" W × 6'10\" H) into a freight elevator. No dock. The stage is "
                    "on the second floor."),
        ("Power", "One 400A 3-phase disconnect, backstage stage right. No shore power; generators are OK."),
        ("House console", "DiGiCo Quantum 225 with Pulse and a DQ Stage rack (48 inputs), in the house right balcony."),
        ("PA", "Sound Bridge line array (main and balcony hangs, L/R), dual 18\" subs and 4 front fills."),
        ("Monitors", "8 wedges and 4 in-ear systems, mixed from FOH. A Yamaha CL3 monitor console and split are available; a fee may apply."),
        ("Lighting", "ETC Ion Element 40 running an LED rep rig with Source 4s, 8 moving heads and 4 truss towers."),
        ("Video", "10,000-lumen laser projector and a 20' motorized 16:9 screen, fed by HDMI from backstage."),
        ("Sound limit", "95 dB LAeq,6min (A-weighted, 6-minute average) at balcony center, per local ordinance."),
        ("Building curfew", "90 minutes after the scheduled end of the show. All touring personnel, artists and gear must be out of the building by then."),
    ], label_w=1.35)

    photos(doc, [(13, "The room from the stage, with the house FOH, house lighting, touring lighting and touring FOH positions marked")], height=1.62)


def stage(doc):
    h1(doc, "Stage", page_break=True)
    spec_table(doc, [
        ("Shape", "Trapezoidal."),
        ("Width", "22'5\" at the back wall; 37'4\" at the widest point."),
        ("Depth", "22'5\" at the deepest point (center stage)."),
        ("Height", "11'5\" at the back-wall ledge; 18' at center stage; 23' at downstage center."),
        ("Shell", "Permanent orchestra shell. The ornate upstage wall is 12' high and scallops to 20'+ at the plaster line."),
        ("Rigging", "No soft goods. No overstage rigging for backdrops, soft goods or lighting."),
        ("Entrances", ["Stage left and stage right doorways, 2'7\" W × 6'10\" H.",
                       "A 5'11\" piano door on the upstage right wall for load-in."]),
        ("House to stage", "Steps from the house to the stage, with or without a railing, can be added in the aisles."),
        ("Building", "Opened 1908. Listed on the National Register of Historic Places in 1978."),
    ])
    photos(doc, [(3, "Stage plan with dimensions")], height=2.9)
    photos(doc, [(1, "The stage from the orchestra"), (2, "Trim and ledge heights on the upstage wall")], height=2.1)


def load_in(doc):
    h1(doc, "Load-in, Parking & Power")
    spec_table(doc, [
        ("Loading door", "On Grant Street at street level. Standard doorway, 2'7\" W × 6'10\" H. Flat push; no loading dock."),
        ("Freight elevator", "Just inside the loading door. 5' W × 7'8\" L × 7' H."),
        ("Stage level", "The stage is on the second floor."),
        ("Parking", "Details come with your advance."),
        ("Power", "One 400A 3-phase disconnect, backstage stage right."),
        ("Shore power", "None. Generators are OK."),
    ])
    photos(doc, [(18, "Loading door on Grant Street"), (17, "Freight elevator, just inside the loading door"),
                 (16, "Inside the freight elevator")], height=2.2)


def audio(doc):
    h1(doc, "Audio")
    h2(doc, "PA System")
    grid_table(doc, ["Qty", "Model", "Position"], [
        ("3 per hang", "Sound Bridge Xyon 7108XY line array", "Main and balcony hangs, L/R"),
        ("2 per side", "Sound Bridge 7218SWX dual 18\" subwoofer", "Ground-stacked, L/R"),
        ("4", "Sound Bridge 3208HTH front fill", "Mono"),
    ], [1.0, 3.3, CONTENT_W - 4.3])

    h2(doc, "Front of House")
    spec_table(doc, [
        ("House console", "DiGiCo Quantum 225 with Pulse."),
        ("Stage rack", "DiGiCo DQ Stage rack, 48 inputs."),
        ("Position", "House right balcony, with a full view of the stage and no booth or overhang. The house console is in a fixed position and does not move."),
    ])
    photos(doc, [(4, "House FOH position, house right balcony (red box)"), (5, "View from the house FOH console")], height=1.95)

    h2(doc, "Touring FOH")
    spec_table(doc, [
        ("Routing", "Guest consoles run through the house console. We can accept left, right, sub, fill and assisted listening "
                    "feeds, or any combination of them."),
        ("Position", "Rear center of the orchestra (Box M), with a full view of the stage and no booth or overhang."),
        ("Cabling", "No drylines at this position; PA drivelines are stage left. The run to the stage box is about 150'. "
                    "Tours must bring all cables needed to connect the touring console to any stage racks or devices."),
    ])
    photos(doc, [(6, "Touring FOH position, Box M (red box)"), (7, "View from the touring FOH position")], height=1.95)

    h2(doc, "Monitors")
    spec_table(doc, [
        ("Wedges", "8 Sound Bridge 7122W 12\" stage monitors."),
        ("In-ears", "4 Sennheiser EW300 in-ear systems."),
        ("Monitor mixes", "Run from FOH on the house console unless a monitor console is in use."),
        ("Monitor console", "Yamaha CL3, 48 inputs. A fee may apply."),
        ("Split", "40-channel transformer-isolated split. It's required when using the house monitor console, and it "
                  "can also feed a touring monitor console. A fee may apply."),
        ("Monitor position", "Offstage left, through a 2'7\" W × 6'10\" H doorway."),
    ])
    photos(doc, [(9, "Monitor position, offstage left"), (8, "View from the monitor position")], height=2.1)

    h2(doc, "Sound Level Limit")
    body(doc, "Local ordinance limits Memorial Hall to 95 dB LAeq,6min: the A-weighted equivalent level averaged over "
              "6 minutes, measured at balcony center.")

    h2(doc, "Recording & Archiving")
    body(doc, "We offer recording and archiving packages. Package options are available on request.")

    h2(doc, "Microphones & DI")
    grid_table(doc, ["Qty", "Model", "Type", "Note"], [
        (8, "Shure SM58", "Dynamic, vocal", ""),
        (8, "Shure SM57", "Dynamic, instrument", ""),
        (1, "Shure Beta 57A", "Dynamic, instrument", ""),
        (1, "Shure Beta 52A", "Dynamic, kick", ""),
        (1, "Shure Beta 91A", "Boundary, kick", ""),
        (2, "Sennheiser e609", "Dynamic, guitar cab", ""),
        (4, "Sennheiser e604", "Dynamic, drum clip", ""),
        (2, "Shure SM81", "Small-diaphragm condenser", ""),
        (2, "Shure Beta 27", "Large-diaphragm condenser", ""),
        (2, "AKG C414", "Large-diaphragm condenser", ""),
        (2, "Telefunken M60", "Small-diaphragm condenser", ""),
        (4, "Shure ULX-D2 / Beta 58A", "Wireless handheld", ""),
        (2, "Shure SLX / SM58", "Wireless handheld", ""),
        (3, "Lavalier mic", "Lavalier", "May incur a fee"),
        (5, "Countryman E6", "Headset", "May incur a fee"),
        (10, "Radial JDI", "Passive DI", ""),
    ], [0.6, 2.2, 2.4, CONTENT_W - 5.2])
    body(doc, "Specific mic requests are possible with advance notice, but additional fees may apply.")


def lighting(doc):
    h1(doc, "Lighting", page_break=True)
    spec_table(doc, [
        ("House console", "ETC Ion Element 40, at the rear of the upper balcony, center."),
        ("House lights", "Our console and crew always run the house lights."),
        ("Tour console", "Touring lighting consoles go in the house left balcony."),
        ("Ground package", "Tours bringing a ground package must supply 150' of cable to reach it."),
        ("Paperwork", "The house rep plot and other lighting paperwork are available on request."),
    ])
    h2(doc, "Inventory")
    grid_table(doc, ["Qty", "Fixture", "Note"], [
        (14, "ETC Source 4, 14° and 19°", ""),
        (10, "ETC Lustre LED, zoom", ""),
        (2, "Robin DLX moving head", ""),
        (6, "LED wash/zoom moving head, 36 × 18W", ""),
        (14, "LED overhead track wash", ""),
        (8, "LED uplight", ""),
        (4, "Elation Fuze 120 LED wash", ""),
        (2, "1 × 1 × 5 stand-alone truss, 4 LED flat pars each", "Mounted"),
        (2, "1 × 1 × 8 stand-alone truss, 4 LED flat pars each", "Mounted"),
        (1, "Radiant water-based hazer", "Additional fee"),
        (2, "Spotlight", "Additional fee"),
    ], [0.6, 4.4, CONTENT_W - 5.0])
    photos(doc, [(10, "View from the house lighting console"),
                 (11, "Touring lighting position in the house left balcony (red box), shot from house right"),
                 (12, "View from the touring lighting position")], height=1.7)


def video(doc):
    h1(doc, "Video", page_break=True)
    spec_table(doc, [
        ("Projector", "Panasonic PT-RZ970BU laser projector, 10,000 lumens, WUXGA."),
        ("Screen", "20' motorized screen, 135\" × 240\", 16:9. It can be partly lowered, with the picture resized to fit."),
        ("Position", "About 13' from the back wall and 9' upstage of the lip. The screen can't be rehung."),
        ("Input", "HDMI from backstage."),
    ])
    photos(doc, [(14, "The screen fully lowered, seen from the orchestra"), (15, "The screen from backstage, stage left")],
           height=2.1)


def equipment(doc):
    h1(doc, "Backline & Stage Equipment")
    body(doc, "The house has no backline instruments or amps, but we can arrange backline rentals.", after=6)
    h2(doc, "Piano")
    body(doc, "Yamaha C3 PE grand (6'1\") with Helpinstill pickups.", after=4)
    h2(doc, "Stands & Seating")
    grid_table(doc, ["Qty", "Item", "Note"], [
        (10, "Tall tripod boom stand", ""),
        (2, "Medium tripod boom stand", ""),
        (6, "Short tripod boom stand", ""),
        (2, "Tall round-base boom stand", ""),
        (5, "Tall round-base straight stand", ""),
        (4, "Short tabletop mic stand", ""),
        (11, "Guitar stand", ""),
        (25, "Music stand", ""),
        (1, "Conductor's music stand", ""),
        (25, "Clip light", "May incur a fee"),
        (20, "Wenger orchestra chair", ""),
        (6, "Backless barstool, unpadded", ""),
    ], [0.6, 4.4, CONTENT_W - 5.0])
    h2(doc, "Risers & Staging")
    grid_table(doc, ["Qty", "Item", "Note"], [
        (4, "Wenger riser, 4' × 8' × 12\"", "Black skirting"),
        (6, "Wenger 3-level choral riser", ""),
        (1, "Conductor's podium", ""),
        ("", "Pipe and drape", "Covers the back wall of the stage"),
        (1, "Large wooden podium", "Gooseneck mic"),
        (1, "Small black podium", "Gooseneck mic"),
        ("", "Drum rugs", ""),
    ], [0.6, 4.4, CONTENT_W - 5.0])


def dressing(doc):
    h1(doc, "Dressing Rooms & Amenities", page_break=True)
    body(doc, "Five dressing rooms, all lockable and climate-controlled.", after=5)
    grid_table(doc, ["Room", "Level", "Capacity", "Features"], [
        ("A", "Stage level, stage right", "2–3", "Vanity mirror and counter; private toilet"),
        ("B", "Stage level, stage left", "2–3", "Vanity mirror and counter; private toilet (identical to A)"),
        ("C", "Balcony level, stage right", "5", "Mirrors, tables, chairs and couches; no private toilet"),
        ("D", "Balcony level, stage left", "5", "Mirrors, tables, chairs and couches; no private toilet"),
        ("Studio", "Basement", "Large group", "Multipurpose room with mirrors, tables and chairs; restroom shared with venue staff"),
    ], [0.7, 1.75, 0.95, CONTENT_W - 3.4])
    body(doc, "Black stage towels are available. There are no showers or laundry facilities on site.", after=6)
    photos(doc, [(19, "Stage-level dressing room (A and B match)"),
                 (20, "Stage-level dressing room, vanity and counter"), (23, "The Studio, basement level")], height=2.3)
    photos(doc, [(21, "Dressing Room C, balcony level, stage right"), (22, "Dressing Room D, balcony level, stage left")],
           height=2.3)


def terms(doc):
    h1(doc, "Additional Terms")
    items = [
        ("Building curfew", "90 minutes after the scheduled end of the show. All touring personnel, artists and gear must be out of the building by then."),
        ("Left behind", "Tell us right away if you leave something behind. Anything left in the building may be "
                        "thrown out."),
        ("Radius", "No other shows within 90 miles of Cincinnati for 60 days before or after your date."),
        ("Merchandise", "The house takes no cut of merch sales. Bring your own POS system and bank. We can help arrange a "
                        "seller if you need one."),
        ("Security", "We may have security on site, but you're responsible for your own gear and belongings."),
    ]
    spec_table(doc, items, label_w=1.35)


def build():
    doc = Document()
    setup(doc)
    page_one(doc)
    stage(doc)
    load_in(doc)
    audio(doc)
    lighting(doc)
    video(doc)
    equipment(doc)
    dressing(doc)
    terms(doc)
    # drop a trailing empty paragraph that would spill a blank page
    last = doc.paragraphs[-1]
    if not last.text:
        last._p.getparent().remove(last._p)
    doc.core_properties.title = "Memorial Hall Technical Specifications"
    doc.core_properties.author = "3CDC Production"
    doc.core_properties.subject = VERSION
    doc.save(OUT + ".docx")
    return OUT + ".docx"


def build_pdf(docx_path):
    pdf = docx_path[:-5] + ".pdf"
    script = f'''
tell application "Microsoft Word"
    set d to open file name (POSIX file "{docx_path}" as text)
    save as d file name (POSIX file "{pdf}" as text) file format format PDF
    close d saving no
end tell'''
    subprocess.run(["osascript", "-e", script], check=True)
    return pdf


if __name__ == "__main__":
    p = build()
    print(p)
    print(build_pdf(p))
